# database.py
import sqlite3
from datetime import datetime, timedelta
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

DB_NAME = 'equipment.db'

def get_connection():
    return sqlite3.connect(DB_NAME)

def init_db():
    with get_connection() as conn:
        cursor = conn.cursor()

        # СОЗДАНИЕ ТАБЛИЦ
        cursor.execute('''CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, password TEXT NOT NULL, role TEXT CHECK(role IN ('admin', 'it_user', 'mol')), full_name TEXT, department TEXT, email TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP)''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS equipment_types (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, category TEXT, depreciation_rate REAL)''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS equipment (id INTEGER PRIMARY KEY AUTOINCREMENT, inventory_number TEXT UNIQUE NOT NULL, serial_number TEXT, name TEXT NOT NULL, type_id INTEGER, status TEXT DEFAULT 'in_stock', location TEXT, responsible_id INTEGER, purchase_date TEXT, warranty_end TEXT, cost REAL, qr_code TEXT, next_maintenance_date TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP, updated_at TEXT DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY (type_id) REFERENCES equipment_types(id), FOREIGN KEY (responsible_id) REFERENCES users(id))''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS movements (id INTEGER PRIMARY KEY AUTOINCREMENT, equipment_id INTEGER NOT NULL, from_mol_id INTEGER, to_mol_id INTEGER, from_location TEXT, to_location TEXT, date TEXT NOT NULL, document_number TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY (equipment_id) REFERENCES equipment(id), FOREIGN KEY (from_mol_id) REFERENCES users(id), FOREIGN KEY (to_mol_id) REFERENCES users(id))''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS repairs (id INTEGER PRIMARY KEY AUTOINCREMENT, equipment_id INTEGER NOT NULL, repair_date TEXT NOT NULL, description TEXT NOT NULL, cost REAL, FOREIGN KEY (equipment_id) REFERENCES equipment(id) ON DELETE CASCADE)''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS licenses (id INTEGER PRIMARY KEY AUTOINCREMENT, equipment_id INTEGER NOT NULL, software_name TEXT NOT NULL, license_key TEXT, expiration_date TEXT NOT NULL, FOREIGN KEY (equipment_id) REFERENCES equipment(id) ON DELETE CASCADE)''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS consumables (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, compatible_models TEXT NOT NULL, quantity INTEGER DEFAULT 0, min_quantity INTEGER DEFAULT 5, unit TEXT DEFAULT 'шт', price REAL, last_updated TEXT DEFAULT CURRENT_TIMESTAMP)''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS compatibility (id INTEGER PRIMARY KEY AUTOINCREMENT, equipment_type_id INTEGER NOT NULL, component_type TEXT NOT NULL, compatible_part TEXT NOT NULL, description TEXT, FOREIGN KEY (equipment_type_id) REFERENCES equipment_types(id))''')
        cursor.execute('''CREATE TABLE IF NOT EXISTS consumables_usage (id INTEGER PRIMARY KEY AUTOINCREMENT, consumable_id INTEGER NOT NULL, equipment_id INTEGER, quantity INTEGER NOT NULL, usage_date TEXT DEFAULT CURRENT_TIMESTAMP, reason TEXT, FOREIGN KEY (consumable_id) REFERENCES consumables(id), FOREIGN KEY (equipment_id) REFERENCES equipment(id))''')

        # НАЧАЛЬНЫЕ ДАННЫЕ
        cursor.execute("SELECT COUNT(*) FROM users")
        if cursor.fetchone()[0] == 0:
            users = [('admin', 'admin123', 'admin', 'Администратор', 'IT-отдел', 'admin@company.ru'), ('ivanov', 'pass123', 'mol', 'Иванов Иван Иванович', 'Бухгалтерия', 'ivanov@company.ru'), ('petrova', 'pass123', 'mol', 'Петрова Анна Сергеевна', 'Отдел продаж', 'petrova@company.ru'), ('sidorov', 'pass123', 'it_user', 'Сидоров Дмитрий Николаевич', 'IT-отдел', 'sidorov@company.ru'), ('kozlov', 'pass123', 'mol', 'Козлов Михаил Андреевич', 'Отдел разработки', 'kozlov@company.ru')]
            cursor.executemany('INSERT INTO users (username, password, role, full_name, department, email) VALUES (?, ?, ?, ?, ?, ?)', users)
            print("  [OK] Добавлены пользователи")

        cursor.execute("SELECT COUNT(*) FROM equipment_types")
        if cursor.fetchone()[0] == 0:
            types = [('ПК', 'Компьютеры', 20.0), ('Ноутбук', 'Компьютеры', 25.0), ('МФУ/Принтер', 'Периферия', 15.0), ('Сетевое оборудование', 'Сетевое', 20.0), ('Сервер', 'Сетевое', 10.0), ('Монитор', 'Периферия', 20.0), ('ИБП', 'Электрооборудование', 15.0)]
            cursor.executemany('INSERT INTO equipment_types (name, category, depreciation_rate) VALUES (?, ?, ?)', types)
            print("  [OK] Добавлены типы оборудования")

        cursor.execute("SELECT COUNT(*) FROM equipment")
        if cursor.fetchone()[0] == 0:
            equipment = [
                ('ИН-001', 'SN100001', 'Системный блок Dell OptiPlex 3080', 1, 'in_use', 'Каб. 101', 2, '2024-03-15', '2026-03-15', 45000),
                ('ИН-002', 'SN100002', 'Ноутбук Lenovo ThinkPad X1 Carbon', 2, 'in_use', 'Каб. 102', 3, '2024-01-10', '2026-01-10', 85000),
                ('ИН-003', 'SN100003', 'МФУ HP LaserJet M404dn', 3, 'repair', 'IT-отдел', 4, '2023-08-20', '2025-08-20', 32000),
                ('ИН-004', 'SN100004', 'Коммутатор Cisco Catalyst 2960', 4, 'in_use', 'Серверная', 4, '2024-05-05', '2026-05-05', 28000),
                ('ИН-005', 'SN100005', 'Монитор Samsung Odyssey G5 27"', 6, 'in_stock', 'Склад №2', None, '2024-12-01', '2026-12-01', 22000),
                ('ИН-006', 'SN100006', 'ИБП APC Smart-UPS 1500', 7, 'in_use', 'Каб. 205', 5, '2023-02-15', '2025-02-15', 38000),
                ('ИН-007', 'SN100007', 'Ноутбук Apple MacBook Air 13" M2', 2, 'in_stock', 'Склад', None, '2025-01-20', '2027-01-20', 120000),
                ('ИН-008', 'SN100008', 'Сервер HP ProLiant DL380 Gen10', 5, 'in_use', 'Серверная', 4, '2023-11-10', '2026-11-10', 320000),
                ('ИН-009', 'SN100009', 'Принтер Brother DCP-L5500DN', 3, 'in_use', 'Бухгалтерия', 2, '2024-06-15', '2026-06-15', 35000),
                ('ИН-010', 'SN100010', 'Монитор LG UltraGear 24"', 6, 'in_stock', 'Склад №1', None, '2025-02-10', '2027-02-10', 15000)
            ]
            cursor.executemany('INSERT INTO equipment (inventory_number, serial_number, name, type_id, status, location, responsible_id, purchase_date, warranty_end, cost) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)', equipment)
            print("  [OK] Добавлено оборудование (10 единиц)")

        cursor.execute("SELECT COUNT(*) FROM movements")
        if cursor.fetchone()[0] == 0:
            movements = [(1, None, 2, None, 'Каб. 101', '2024-03-15', 'АКТ-001'), (2, None, 3, None, 'Каб. 102', '2024-01-10', 'АКТ-002'), (3, 4, 4, 'IT-отдел', 'Мастерская', '2024-10-01', 'РЕМ-001')]
            cursor.executemany('INSERT INTO movements (equipment_id, from_mol_id, to_mol_id, from_location, to_location, date, document_number) VALUES (?, ?, ?, ?, ?, ?, ?)', movements)
            print("  [OK] Добавлена история перемещений")

        cursor.execute("SELECT COUNT(*) FROM consumables")
        if cursor.fetchone()[0] == 0:
            consumables = [('Картридж HP 85A', 'HP LaserJet M404dn', 15, 5, 'шт', 1200), ('Тонер Brother TN-1050', 'Brother DCP-L5500DN', 8, 3, 'шт', 1500), ('Бумага А4 500л', 'Универсальный', 50, 10, 'пачка', 350)]
            cursor.executemany('INSERT INTO consumables (name, compatible_models, quantity, min_quantity, unit, price) VALUES (?, ?, ?, ?, ?, ?)', consumables)
            print("  [OK] Добавлены расходные материалы")

        cursor.execute("SELECT COUNT(*) FROM compatibility")
        if cursor.fetchone()[0] == 0:
            compatibility = [(1, 'ОЗУ', 'DDR4 8GB 2666MHz', 'Kingston / Samsung / Crucial'), (2, 'ОЗУ', 'DDR4 8GB 3200MHz SO-DIMM', 'Kingston / Samsung / Crucial')]
            cursor.executemany('INSERT INTO compatibility (equipment_type_id, component_type, compatible_part, description) VALUES (?, ?, ?, ?)', compatibility)
            print("  [OK] Добавлена совместимость комплектующих")

        cursor.execute("SELECT COUNT(*) FROM licenses")
        if cursor.fetchone()[0] == 0:
            licenses = [
                (1, 'Windows 11 Pro', 'WX4-1111-2222-3333', (datetime.now() + timedelta(days=15)).strftime('%Y-%m-%d')),
                (2, 'Microsoft Office 2021', 'OF-4444-5555-6666', (datetime.now() + timedelta(days=45)).strftime('%Y-%m-%d')),
                (3, 'Windows 10 Pro', 'WX4-7777-8888-9999', (datetime.now() - timedelta(days=10)).strftime('%Y-%m-%d')),
            ]
            cursor.executemany('INSERT INTO licenses (equipment_id, software_name, license_key, expiration_date) VALUES (?, ?, ?, ?)', licenses)
            print("  [OK] Добавлены лицензии")

        cursor.execute("SELECT COUNT(*) FROM repairs")
        if cursor.fetchone()[0] == 0:
            repairs = [
                (3, (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d'), 'Замена тонера', 2500),
                (3, (datetime.now() - timedelta(days=15)).strftime('%Y-%m-%d'), 'Прочистка печатающей головки', 0),
                (1, (datetime.now() - timedelta(days=60)).strftime('%Y-%m-%d'), 'Замена блока питания', 3500),
            ]
            cursor.executemany('INSERT INTO repairs (equipment_id, repair_date, description, cost) VALUES (?, ?, ?, ?)', repairs)
            print("  [OK] Добавлена история ремонтов")

        cursor.execute("UPDATE equipment SET next_maintenance_date = date('now', '+5 days') WHERE id = 1 AND next_maintenance_date IS NULL")
        cursor.execute("UPDATE equipment SET next_maintenance_date = date('now', '+12 days') WHERE id = 2 AND next_maintenance_date IS NULL")
        conn.commit()
    print("\n[OK] База данных успешно инициализирована!")

# ОСНОВНЫЕ ОПЕРАЦИИ
def get_statistics():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM equipment")
        total = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(CASE WHEN status = 'in_use' THEN 1 END), COUNT(CASE WHEN status = 'in_stock' THEN 1 END), COUNT(CASE WHEN status = 'repair' THEN 1 END) FROM equipment")
        row = cursor.fetchone()
        return {'total': total, 'in_use': row[0] or 0, 'in_stock': row[1] or 0, 'repair': row[2] or 0}

def get_all_equipment():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT e.id, e.inventory_number, e.serial_number, e.name, e.status, e.location, e.purchase_date, e.cost, e.warranty_end, et.name as type_name, u.full_name as responsible_name FROM equipment e LEFT JOIN equipment_types et ON e.type_id = et.id LEFT JOIN users u ON e.responsible_id = u.id ORDER BY e.id')
        return cursor.fetchall()

def get_equipment_by_id(equipment_id):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT e.id, e.inventory_number, e.serial_number, e.name, e.status, e.location, e.purchase_date, e.cost, e.warranty_end, e.type_id, e.responsible_id, et.name as type_name, u.full_name as responsible_name FROM equipment e LEFT JOIN equipment_types et ON e.type_id = et.id LEFT JOIN users u ON e.responsible_id = u.id WHERE e.id = ?', (equipment_id,))
        return cursor.fetchone()

def add_equipment(inventory_number, name, serial_number='', type_id=None, status='in_stock', location='', responsible_id=None, purchase_date='', warranty_end='', cost=None, qr_code=''):
    with get_connection() as conn:
        cursor = conn.cursor()
        if not purchase_date: purchase_date = datetime.now().strftime('%Y-%m-%d')
        cursor.execute('INSERT INTO equipment (inventory_number, serial_number, name, type_id, status, location, responsible_id, purchase_date, warranty_end, cost, qr_code) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)', (inventory_number, serial_number, name, type_id, status, location, responsible_id, purchase_date, warranty_end, cost, qr_code))
        conn.commit()
        return cursor.lastrowid

def update_qr_code(equipment_id, qr_code):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('UPDATE equipment SET qr_code = ? WHERE id = ?', (qr_code, equipment_id))
        conn.commit()

def update_equipment(equipment_id, inventory_number, name, serial_number='', status='in_stock', location=''):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('UPDATE equipment SET inventory_number=?, name=?, serial_number=?, status=?, location=?, updated_at=CURRENT_TIMESTAMP WHERE id=?', (inventory_number, name, serial_number, status, location, equipment_id))
        conn.commit()

def delete_equipment(equipment_id):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('DELETE FROM equipment WHERE id = ?', (equipment_id,))
        conn.commit()

def search_equipment(search_term):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT id, inventory_number, name, status, location FROM equipment WHERE inventory_number LIKE ? OR name LIKE ?', (f'%{search_term}%', f'%{search_term}%'))
        return cursor.fetchall()

def get_statistics_for_dashboard():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT et.name, COUNT(e.id) FROM equipment_types et LEFT JOIN equipment e ON e.type_id = et.id GROUP BY et.id')
        type_stats = [{'name': row[0], 'count': row[1]} for row in cursor.fetchall()]
        cursor.execute('SELECT strftime("%Y-%m", created_at) as month, COUNT(*) FROM equipment GROUP BY month ORDER BY month DESC LIMIT 6')
        monthly_stats = [{'month': row[0], 'count': row[1]} for row in cursor.fetchall()]
        monthly_stats.reverse()
        cursor.execute('SELECT e.name, COUNT(r.id) as repair_count, e.inventory_number FROM equipment e LEFT JOIN repairs r ON e.id = r.equipment_id GROUP BY e.id ORDER BY repair_count DESC LIMIT 5')
        top_problematic = [{'name': row[0], 'repairs': row[1], 'inv': row[2]} for row in cursor.fetchall()]
        cursor.execute('SELECT name, inventory_number, warranty_end FROM equipment WHERE warranty_end BETWEEN date("now") AND date("now", "+3 months") ORDER BY warranty_end')
        expiring_warranty = [{'name': row[0], 'inv': row[1], 'end': row[2]} for row in cursor.fetchall()]
        return {'by_type': type_stats, 'by_month': monthly_stats, 'top_problematic': top_problematic, 'expiring_warranty': expiring_warranty}

def get_equipment_history(equipment_id):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT m.*, u1.full_name as from_mol_name, u2.full_name as to_mol_name FROM movements m LEFT JOIN users u1 ON m.from_mol_id = u1.id LEFT JOIN users u2 ON m.to_mol_id = u2.id WHERE m.equipment_id = ? ORDER BY m.date DESC', (equipment_id,))
        return cursor.fetchall()

def export_to_excel():
    equipment = get_all_equipment()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Оборудование"
    headers = ['ID', 'Инв. номер', 'Серийный номер', 'Наименование', 'Тип', 'Статус', 'Местоположение', 'МОЛ', 'Дата покупки', 'Гарантия до', 'Стоимость']
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2a5298", end_color="2a5298", fill_type="solid")
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for eq in equipment:
        status_display = {'in_use': 'В эксплуатации', 'in_stock': 'На складе', 'repair': 'В ремонте'}.get(eq[4], eq[4])
        ws.append([eq[0], eq[1], eq[2] or '', eq[3], eq[9] or '', status_display, eq[5] or '', eq[10] or '', eq[6] or '', eq[8] or '', eq[7] or 0])
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_len: max_len = len(str(cell.value))
            except: pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 30)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# ОТЧЕТЫ И АНАЛИТИКА
def add_repair(equipment_id, repair_date, description, cost=None):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('INSERT INTO repairs (equipment_id, repair_date, description, cost) VALUES (?, ?, ?, ?)', (equipment_id, repair_date, description, cost))
        conn.commit()

def get_all_repairs_with_equipment():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT r.*, e.inventory_number, e.name as equipment_name FROM repairs r JOIN equipment e ON r.equipment_id = e.id ORDER BY r.repair_date DESC')
        return cursor.fetchall()

def get_problem_equipment():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT e.id, e.inventory_number, e.name, COUNT(r.id) as repair_count, COALESCE(SUM(r.cost), 0) as total_cost FROM equipment e LEFT JOIN repairs r ON e.id = r.equipment_id GROUP BY e.id ORDER BY repair_count DESC LIMIT 10')
        return cursor.fetchall()

def add_license(equipment_id, software_name, license_key, expiration_date):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('INSERT INTO licenses (equipment_id, software_name, license_key, expiration_date) VALUES (?, ?, ?, ?)', (equipment_id, software_name, license_key, expiration_date))
        conn.commit()

def get_all_licenses_with_equipment():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT l.*, e.inventory_number, e.name as equipment_name FROM licenses l JOIN equipment e ON l.equipment_id = e.id ORDER BY l.expiration_date ASC')
        return cursor.fetchall()

def get_expiring_licenses(days=30):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT l.*, e.inventory_number, e.name as equipment_name FROM licenses l JOIN equipment e ON l.equipment_id = e.id WHERE l.expiration_date BETWEEN date('now') AND date('now', ?) ORDER BY l.expiration_date ASC", (f'+{days} days',))
        return cursor.fetchall()

def get_upcoming_maintenance(days=30):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT e.id, e.inventory_number, e.name, e.location, e.responsible_id, u.full_name as responsible_name, e.next_maintenance_date FROM equipment e LEFT JOIN users u ON e.responsible_id = u.id WHERE e.next_maintenance_date IS NOT NULL AND e.next_maintenance_date BETWEEN date('now') AND date('now', ?) ORDER BY e.next_maintenance_date ASC", (f'+{days} days',))
        return cursor.fetchall()

def get_all_consumables():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT * FROM consumables ORDER BY name')
        return cursor.fetchall()

def add_consumable(name, compatible_models, quantity, min_quantity=5, unit='шт', price=None):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('INSERT INTO consumables (name, compatible_models, quantity, min_quantity, unit, price) VALUES (?, ?, ?, ?, ?, ?)', (name, compatible_models, quantity, min_quantity, unit, price))
        conn.commit()

def use_consumable(consumable_id, equipment_id, quantity, reason=''):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT quantity FROM consumables WHERE id = ?', (consumable_id,))
        current = cursor.fetchone()
        if not current or current[0] < quantity: raise ValueError("Недостаточное количество расходного материала")
        cursor.execute('UPDATE consumables SET quantity = quantity - ?, last_updated = CURRENT_TIMESTAMP WHERE id = ?', (quantity, consumable_id))
        cursor.execute('INSERT INTO consumables_usage (consumable_id, equipment_id, quantity, reason) VALUES (?, ?, ?, ?)', (consumable_id, equipment_id, quantity, reason))
        conn.commit()

def get_low_stock_consumables():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT * FROM consumables WHERE quantity <= min_quantity ORDER BY quantity ASC')
        return cursor.fetchall()

def get_compatible_parts(equipment_type_id):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT * FROM compatibility WHERE equipment_type_id = ?', (equipment_type_id,))
        return cursor.fetchall()

def get_most_problematic_equipment(limit=10):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT e.id, e.inventory_number, e.name, COUNT(r.id) as repair_count, COALESCE(SUM(r.cost), 0) as total_repair_cost FROM equipment e LEFT JOIN repairs r ON e.id = r.equipment_id GROUP BY e.id ORDER BY repair_count DESC, total_repair_cost DESC LIMIT ?', (limit,))
        return cursor.fetchall()

def get_oldest_equipment(limit=10):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT id, inventory_number, name, purchase_date, julianday("now") - julianday(purchase_date) as days_old, cost, status FROM equipment WHERE purchase_date IS NOT NULL ORDER BY purchase_date ASC LIMIT ?', (limit,))
        return cursor.fetchall()

def get_requires_maintenance(limit=10):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT e.id, e.inventory_number, e.name, e.location, e.next_maintenance_date, u.full_name as responsible_name FROM equipment e LEFT JOIN users u ON e.responsible_id = u.id WHERE e.next_maintenance_date IS NOT NULL AND e.next_maintenance_date BETWEEN date('now') AND date('now', '+30 days') ORDER BY e.next_maintenance_date ASC LIMIT ?", (limit,))
        return cursor.fetchall()

def get_expiring_warranty_equipment(limit=10):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT id, inventory_number, name, warranty_end, cost, status FROM equipment WHERE warranty_end IS NOT NULL AND warranty_end BETWEEN date("now") AND date("now", "+90 days") ORDER BY warranty_end ASC LIMIT ?', (limit,))
        return cursor.fetchall()

def get_most_moved_equipment(limit=10):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT e.id, e.inventory_number, e.name, COUNT(m.id) as move_count, e.location, u.full_name as responsible_name FROM equipment e LEFT JOIN movements m ON e.id = m.equipment_id LEFT JOIN users u ON e.responsible_id = u.id GROUP BY e.id ORDER BY move_count DESC LIMIT ?', (limit,))
        return cursor.fetchall()

def get_most_expensive_equipment(limit=10):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT id, inventory_number, name, cost, status, location FROM equipment WHERE cost IS NOT NULL ORDER BY cost DESC LIMIT ?', (limit,))
        return cursor.fetchall()

def get_all_types():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT id, name FROM equipment_types ORDER BY name')
        return cursor.fetchall()